
from fpdf import FPDF
from fpdf.fonts import FontFace
from database import get_expenses

def get_filtered_expenses(chat_type, identifier):
    data = get_expenses()
    if chat_type == 'private':
        return [d for d in data if d.get('user_id') == identifier and d.get('group_name') == 'Shaxsiy']
    else:
        return [d for d in data if d.get('group_name') == identifier]

import openpyxl

def export_to_excel(chat_type, identifier, filename="xarajatlar.xlsx"):
    expenses = get_filtered_expenses(chat_type, identifier)
    if not expenses:
        return None
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Xarajatlar"
    
    headers = ["Foydalanuvchi ID", "Summa", "Mahsulotlar/Tavsif", "Do'kon nomi", "Guruh/Chat", "Sana"]
    ws.append(headers)
    
    for exp in expenses:
        date_str = str(exp.get('date', ''))[:16].replace('T', ' ')
        desc = exp.get('description', '')
        if isinstance(desc, list):
            desc = "\n".join(str(d) for d in desc)
        else:
            desc = str(desc)
            
        ws.append([
            exp.get('user_id', ''),
            exp.get('amount', 0),
            desc,
            exp.get('shop_name', "Noma'lum"),
            exp.get('group_name', "Shaxsiy"),
            date_str
        ])
        
    wb.save(filename)
    return filename

class PDF(FPDF):
    def header(self):
        try:
            self.add_font("ArialCyr", "", "arial.ttf")
            self.add_font("ArialCyr", "B", "arialbd.ttf")
            self.set_font("ArialCyr", "B", 16)
        except:
            self.set_font("helvetica", "B", 16)
        self.cell(0, 10, "Xarajatlar Hisoboti", align="C")
        self.ln(15)

def export_to_pdf(chat_type, identifier, filename="hisobot.pdf"):
    expenses = get_filtered_expenses(chat_type, identifier)
    if not expenses:
        return None
    
    pdf = PDF()
    pdf.add_page()
    
    # Calculate Total
    total_amount = sum(float(exp.get('amount', 0)) for exp in expenses)
    
    try:
        pdf.set_font("ArialCyr", "B", 12)
    except:
        pdf.set_font("helvetica", "B", 12)
        
    group_display = str(identifier if chat_type != 'private' else 'Shaxsiy xarajatlar')
    pdf.cell(0, 8, f"Guruh / Chat: {group_display}", align="L")
    pdf.ln(8)
    
    formatted_total = f"{int(total_amount):,} so'm".replace(',', ' ')
    pdf.cell(0, 8, f"Jami xarajatlar summasi: {formatted_total}", align="L")
    pdf.ln(12)
    
    try:
        pdf.set_font("ArialCyr", size=9)
    except:
        pdf.set_font("helvetica", size=9)
        
    # Define table styles
    head_style = FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=(41, 128, 185))
    
    with pdf.table(col_widths=(25, 30, 85, 30), text_align=("C", "L", "L", "R")) as table:
        # Header
        row = table.row()
        row.cell("Sana", style=head_style)
        row.cell("Do'kon", style=head_style)
        row.cell("Mahsulotlar", style=head_style)
        row.cell("Summa", style=head_style)
        
        # Data Rows
        for exp in expenses:
            row = table.row()
            
            date = str(exp.get('date', ''))[:16].replace('T', '\n')
            shop = str(exp.get('shop_name', "Noma'lum"))
            amount_val = float(exp.get('amount', 0))
            amt_str = f"{int(amount_val):,} so'm".replace(',', ' ')
            
            desc = exp.get('description', '')
            if isinstance(desc, list):
                desc = "\n".join(str(d) for d in desc)
            else:
                desc = str(desc)
            
            row.cell(date)
            row.cell(shop)
            row.cell(desc)
            row.cell(amt_str)
            
    pdf.output(filename)
    return filename
